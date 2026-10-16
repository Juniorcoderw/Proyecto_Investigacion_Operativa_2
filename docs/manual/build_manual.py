from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
MANUAL_DIR = ROOT / "docs" / "manual"
SHOTS = MANUAL_DIR / "recursos"
OUT = MANUAL_DIR / "Manual_de_usuario_Investigacion_Operativa_II.docx"

NAVY = "0C2340"
GOLD = "D5AD45"
GREEN = "1B4D43"
BURGUNDY = "4A2238"
CORAL = "B9573F"
INK = "1B2430"
BODY = "3F4854"
MUTED = "6D7480"
LIGHT = "F5F3ED"
PALE_BLUE = "EEF3F7"
PALE_GOLD = "FBF6E8"
BORDER = "D9D9D9"
WHITE = "FFFFFF"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in edges:
            continue
        data = edges[edge]
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in data:
                element.set(qn("w:{}".format(key)), str(data[key]))


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep(paragraph, keep_next=False, keep_lines=True):
    paragraph.paragraph_format.keep_with_next = keep_next
    paragraph.paragraph_format.keep_together = keep_lines


def set_run_font(run, name="Aptos", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_hyperlink(paragraph, text, url, color=NAVY, underline=True):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Aptos")
    r_fonts.set(qn("w:hAnsi"), "Aptos")
    r_pr.append(r_fonts)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def add_image_border(inline_shape, color=NAVY, width="12700"):
    sp_pr = inline_shape._inline.graphic.graphicData.pic.spPr
    line = sp_pr.find(qn("a:ln"))
    if line is None:
        line = OxmlElement("a:ln")
        sp_pr.append(line)
    line.set("w", width)
    solid = OxmlElement("a:solidFill")
    srgb = OxmlElement("a:srgbClr")
    srgb.set("val", color)
    solid.append(srgb)
    line.append(solid)


def add_picture(doc, filename, caption, alt_text, width=6.65):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    set_keep(p, keep_next=True)
    shape = p.add_run().add_picture(str(SHOTS / filename), width=Inches(width))
    add_image_border(shape)
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", alt_text)
    doc_pr.set("title", caption)

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(7)
    set_keep(cap)
    r = cap.add_run(caption)
    set_run_font(r, size=9, color=MUTED, italic=True)
    return shape


def add_section_heading(doc, number, title, intro=None, accent=NAVY):
    tag = doc.add_paragraph()
    tag.paragraph_format.space_before = Pt(0)
    tag.paragraph_format.space_after = Pt(3)
    set_keep(tag, keep_next=True)
    r = tag.add_run(f"{number:02d}  GUÍA DE USO")
    set_run_font(r, size=9, color=accent, bold=True)

    h = doc.add_heading(title, level=1)
    h.paragraph_format.space_after = Pt(6)
    set_keep(h, keep_next=True)
    if intro:
        p = doc.add_paragraph(intro)
        p.paragraph_format.space_after = Pt(8)
        set_keep(p)


def add_steps(doc, steps, accent=NAVY):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(0.48)
    table.columns[1].width = Inches(6.05)
    for idx, text in enumerate(steps, 1):
        row = table.add_row()
        row.cells[0].width = Inches(0.48)
        row.cells[1].width = Inches(6.05)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            set_cell_border(
                cell,
                top={"val": "single", "sz": 5, "color": BORDER},
                bottom={"val": "single", "sz": 5, "color": BORDER},
                left={"val": "single", "sz": 5, "color": BORDER},
                right={"val": "single", "sz": 5, "color": BORDER},
            )
        set_cell_fill(row.cells[0], accent)
        set_cell_fill(row.cells[1], "FFFFFF" if idx % 2 else PALE_BLUE)
        p_num = row.cells[0].paragraphs[0]
        p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_num = p_num.add_run(str(idx))
        set_run_font(r_num, size=10, color=WHITE, bold=True)
        p_text = row.cells[1].paragraphs[0]
        p_text.paragraph_format.space_after = Pt(0)
        r_text = p_text.add_run(text)
        set_run_font(r_text, size=10.5, color=BODY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_module_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.4, 2.25, 2.9]
    headers = ["Módulo", "Se usa para", "Resultado principal"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell, 100, 100, 100, 100)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=9.5, color=WHITE, bold=True)
    set_repeat_table_header(table.rows[0])
    rows = [
        ("PERT CPM", "Planificar actividades con precedencias", "Red AON ruta crítica Gantt y tiempos"),
        ("Inventarios", "Definir lotes pedidos o producción", "Cantidad óptima punto de reorden y costos"),
        ("Decisiones", "Comparar alternativas bajo incertidumbre", "Criterios VME EVPI y estrategia del árbol"),
    ]
    for ridx, values in enumerate(rows, 1):
        row = table.add_row()
        for i, text in enumerate(values):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_fill(cell, WHITE if ridx % 2 else PALE_BLUE)
            set_cell_margins(cell, 105, 105, 105, 105)
            set_cell_border(
                cell,
                top={"val": "single", "sz": 5, "color": BORDER},
                bottom={"val": "single", "sz": 5, "color": BORDER},
                left={"val": "single", "sz": 5, "color": BORDER},
                right={"val": "single", "sz": 5, "color": BORDER},
            )
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            set_run_font(r, size=9.5, color=BODY, bold=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_lead(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(label + " ")
    set_run_font(r, size=10.5, color=INK, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=BODY)


def page_break(doc):
    doc.add_page_break()


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for style_name, size in (("Heading 1", 21), ("Heading 2", 14), ("Heading 3", 11.5)):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run("Investigación Operativa II   •   io2modelos.vercel.app   •   ")
    set_run_font(r, size=8, color=MUTED)
    add_page_number(p)


def build():
    doc = Document()
    configure_document(doc)

    # Portada
    logos = doc.add_paragraph()
    logos.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logos.paragraph_format.space_after = Pt(10)
    r = logos.add_run()
    logo_unjbg = r.add_picture(str(ROOT / "assets" / "institucional" / "logo-unjbg.png"), width=Inches(0.78))
    logo_unjbg._inline.docPr.set("descr", "Escudo de la Universidad Nacional Jorge Basadre Grohmann")
    logo_unjbg._inline.docPr.set("title", "Universidad Nacional Jorge Basadre Grohmann")
    logos.add_run("      ")
    r = logos.add_run()
    logo_esis = r.add_picture(str(ROOT / "assets" / "institucional" / "logo-esis.png"), width=Inches(0.75))
    logo_esis._inline.docPr.set("descr", "Logotipo de Ingeniería en Informática y Sistemas de la UNJBG")
    logo_esis._inline.docPr.set("title", "Ingeniería en Informática y Sistemas")

    t = doc.add_paragraph(style="Title")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(5)
    t.add_run("Manual de usuario de la plataforma de Investigación Operativa II")
    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.paragraph_format.space_after = Pt(12)
    r = st.add_run("PERT CPM Inventarios y Teoría de Decisiones")
    set_run_font(r, name="Aptos Display", size=15, color=INK, bold=True)

    add_picture(
        doc,
        "01-portada.png",
        "Vista principal de la plataforma",
        "Portada de la plataforma de Investigación Operativa II con acceso a PERT CPM Inventarios y Decisiones",
        width=6.45,
    )

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(7)
    meta.paragraph_format.space_after = Pt(3)
    r = meta.add_run("Universidad Nacional Jorge Basadre Grohmann\nEscuela Profesional de Ingeniería en Informática y Sistemas")
    set_run_font(r, size=10.5, color=INK, bold=True)
    meta2 = doc.add_paragraph()
    meta2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta2.add_run("Docente MSc Manuel Yuri Apaza Valencia   •   Tacna 2026")
    set_run_font(r, size=9.5, color=MUTED)

    page_break(doc)

    # Introducción y navegación
    add_section_heading(
        doc,
        1,
        "Acceso y navegación",
        "Este manual explica cómo ingresar datos, ejecutar un cálculo y revisar los resultados de cada módulo. La plataforma funciona en un navegador moderno y conserva los últimos datos en el equipo utilizado.",
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Dirección de acceso  ")
    set_run_font(r, size=10.5, color=INK, bold=True)
    add_hyperlink(p, "Abrir la plataforma", "https://io2modelos.vercel.app/")
    add_picture(
        doc,
        "02-modulos.png",
        "Figura 1  Selector de módulos de la plataforma",
        "Tres tarjetas para abrir los módulos PERT CPM Inventarios y Decisiones",
        width=6.45,
    )
    add_steps(
        doc,
        [
            "Abra Calculadoras o desplácese hasta la sección de módulos.",
            "Elija PERT CPM Inventarios o Decisiones según el problema que desea resolver.",
            "Use un caso precargado para aprender el flujo o reemplace los valores con su ejercicio.",
            "Abra Guía del método cuando necesite revisar símbolos supuestos o fórmulas.",
        ],
    )

    page_break(doc)

    # Mapa de módulos
    add_section_heading(
        doc,
        2,
        "Elección del módulo",
        "Cada calculadora sigue el mismo recorrido general Datos Solución y Análisis. Elija el módulo por el tipo de problema y no por la forma de los datos.",
    )
    add_module_table(doc)
    doc.add_heading("Flujo común", level=2)
    add_steps(
        doc,
        [
            "Revise el modelo seleccionado y la unidad de cada campo antes de escribir.",
            "Cargue un caso de práctica o complete el formulario con los datos del ejercicio.",
            "Pulse el botón principal de cálculo. Los resultados aparecen en el panel derecho.",
            "Cambie entre las pestañas de análisis para ver gráficos desarrollo matemático y comparaciones.",
            "Use Imprimir informe si necesita conservar el resultado o presentarlo junto al desarrollo manual.",
        ],
        accent=GREEN,
    )
    add_lead(doc, "Guardado local", "Los datos se almacenan en el navegador del mismo equipo. No se sincronizan entre computadoras ni entre navegadores distintos.")
    add_lead(doc, "Casos de práctica", "Cargar un ejemplo reemplaza el contenido actual del formulario. Si necesita conservarlo exporte o imprima el resultado antes de cargar otro caso.")

    page_break(doc)

    # PERT ingreso
    add_section_heading(
        doc,
        3,
        "Ingreso de datos en PERT CPM",
        "Use CPM cuando cada actividad tenga una duración conocida. Use PERT cuando disponga de tiempos optimista más probable y pesimista.",
        accent=CORAL,
    )
    add_picture(
        doc,
        "03-pert-ingreso.png",
        "Figura 2  Formulario de actividades PERT CPM",
        "Formulario PERT CPM con selector de método tabla de actividades casos de práctica y botón Calcular proyecto",
        width=6.45,
    )
    add_steps(
        doc,
        [
            "Seleccione CPM determinístico o PERT de tres estimaciones.",
            "Cargue un caso de práctica o agregue filas con el botón Fila.",
            "Escriba un identificador una descripción las predecesoras y la duración. Separe varias predecesoras con comas como B, C.",
            "También puede pegar una tabla desde Excel o importar un archivo CSV.",
            "Pulse Calcular proyecto para generar la red la ruta crítica el Gantt y la solución.",
        ],
        accent=CORAL,
    )

    page_break(doc)

    # PERT resultados
    add_section_heading(
        doc,
        4,
        "Lectura de resultados en PERT CPM",
        "La red muestra una actividad por nodo. Las ramas críticas se resaltan y las actividades con holgura se distinguen con un segundo color.",
        accent=CORAL,
    )
    add_picture(
        doc,
        "04-pert-resultado.png",
        "Figura 3  Red del proyecto y resumen de la ruta crítica",
        "Red AON con nodos de actividades ruta crítica controles de zoom y tarjetas de duración",
        width=6.45,
    )
    add_steps(
        doc,
        [
            "En Red del proyecto use menos Ajustar y más para cambiar el nivel de zoom.",
            "Haga clic en cualquier nodo para abrir el cálculo de esa actividad.",
            "Revise Rutas críticas y los indicadores de duración total cantidad de rutas y actividades críticas.",
            "Abra Gantt para comprobar el calendario y Solución para revisar el cálculo en orden.",
            "Pulse Imprimir informe cuando el análisis esté listo.",
        ],
        accent=CORAL,
    )

    page_break(doc)

    # PERT solución y guía
    add_section_heading(
        doc,
        5,
        "Desarrollo y guía del método",
        "La vista Solución presenta las operaciones por actividad para facilitar la comparación con un desarrollo manual.",
        accent=CORAL,
    )
    add_picture(
        doc,
        "05-pert-solucion.png",
        "Figura 4  Desarrollo paso a paso del proyecto",
        "Vista Solución de PERT CPM con recorrido hacia adelante y operaciones por actividad",
        width=6.25,
    )
    add_steps(
        doc,
        [
            "Abra Solución y despliegue Avanzar por la red Retroceder desde el final o Calcular el margen disponible.",
            "Lea cada fila de izquierda a derecha para identificar los valores usados y el resultado de la actividad.",
            "Use Nodo para relacionar la operación con el cuadro correspondiente del diagrama.",
        ],
        accent=CORAL,
    )

    page_break(doc)
    add_section_heading(
        doc,
        6,
        "Consulta de fundamentos",
        "La guía reúne la teoría básica los símbolos y la formulación matemática sin abandonar el ejercicio actual.",
        accent=CORAL,
    )
    add_picture(
        doc,
        "06-guia-metodo.png",
        "Figura 5  Guía flotante con teoría y modelamiento matemático",
        "Ventana flotante de fundamentos y formulación matemática PERT CPM",
        width=5.9,
    )
    add_steps(
        doc,
        [
            "Pulse Guía del método desde cualquier vista del módulo.",
            "Abra Fundamentos y formulación matemática para revisar el significado de cada valor y su cálculo.",
            "Cierre la ventana con la X. Los datos y el resultado del ejercicio permanecen disponibles.",
        ],
        accent=CORAL,
    )

    page_break(doc)

    # Inventarios ingreso
    add_section_heading(
        doc,
        7,
        "Ingreso de datos en inventarios",
        "El módulo incluye EOQ con demanda constante EPQ con producción gradual y EOQ con faltantes planificados.",
        accent=GREEN,
    )
    add_picture(
        doc,
        "07-inventarios-ingreso.png",
        "Figura 6  Formulario de inventarios",
        "Formulario de inventarios con selección de modelo campos de demanda costos días y moneda",
        width=6.45,
    )
    add_steps(
        doc,
        [
            "Seleccione el modelo que corresponde al ejercicio.",
            "Escriba la demanda el costo por pedido el costo unitario los días laborables y el tiempo de entrega usando el mismo periodo.",
            "Ingrese el costo de mantenimiento por unidad o como porcentaje del precio. Use capacidad máxima solo cuando el ejercicio la proporcione.",
            "Seleccione un caso precargado y pulse Cargar si desea comparar con una práctica.",
            "Pulse Calcular política óptima.",
        ],
        accent=GREEN,
    )

    page_break(doc)

    # Inventarios resultados
    add_section_heading(
        doc,
        8,
        "Lectura de resultados en inventarios",
        "La primera fila resume la política recomendada. El gráfico muestra cuándo realizar el pedido y cómo cambia el inventario con el tiempo.",
        accent=GREEN,
    )
    add_picture(
        doc,
        "08-inventarios-resultado.png",
        "Figura 7  Política óptima y gráfico de inventario",
        "Resultados EOQ con cantidad económica punto de reorden frecuencia tiempo entre pedidos y gráfico de existencias",
        width=6.45,
    )
    add_steps(
        doc,
        [
            "Lea la cantidad económica Q el punto de reorden r los pedidos por año N y el tiempo entre pedidos T.",
            "Use la recomendación escrita para convertir los resultados en una política de pedido.",
            "En Política cambie el número de ciclos para observar más o menos periodos en el gráfico.",
            "Abra Costos para revisar los componentes anuales Desarrollo para comprobar las operaciones y Comparar para evaluar escenarios.",
            "Use Imprimir informe para guardar la política calculada.",
        ],
        accent=GREEN,
    )

    page_break(doc)

    # Tablas de decisión
    add_section_heading(
        doc,
        9,
        "Tablas de decisión",
        "Use esta vista cuando exista una decisión con varias alternativas y estados de la naturaleza definidos en una sola matriz.",
        accent=BURGUNDY,
    )
    add_picture(
        doc,
        "09-decisiones-tabla.png",
        "Figura 8  Matriz de pagos y resultados por criterio",
        "Tabla de decisión con alternativas estados probabilidades y resumen de valor esperado",
        width=6.45,
    )
    add_steps(
        doc,
        [
            "Cargue un caso o escriba el nombre del problema.",
            "Indique si los valores son beneficios para elegir el mayor o costos para elegir el menor.",
            "Ajuste el optimismo de Hurwicz y edite alternativas estados pagos y probabilidades. La suma debe ser 100 por ciento.",
            "Pulse Evaluar decisiones.",
            "Revise la recomendación el mejor valor esperado el EVPI y las pestañas Criterios Arrepentimiento y Desarrollo.",
        ],
        accent=BURGUNDY,
    )

    page_break(doc)

    # Árboles de decisión
    add_section_heading(
        doc,
        10,
        "Árboles de decisión",
        "Use esta vista cuando una decisión conduzca a eventos posteriores. El sistema resuelve el árbol desde los resultados finales hasta el nodo inicial.",
        accent=BURGUNDY,
    )
    add_picture(
        doc,
        "10-arbol-decision.png",
        "Figura 9  Editor y diagrama del árbol de decisión",
        "Árbol de decisión con editor de nodos controles de zoom y estrategia óptima resaltada",
        width=6.45,
    )
    add_steps(
        doc,
        [
            "Cargue un caso o cree un árbol nuevo y defina si desea maximizar beneficios o minimizar costos.",
            "Seleccione un nodo en el diagrama para editar su nombre tipo y ramas.",
            "Use Decisión para elegir una rama Azar para ponderar probabilidades y Resultado final para ingresar un valor terminal.",
            "Pulse Resolver árbol. La estrategia recomendada se resalta en dorado y el valor se muestra cerca de cada nodo.",
            "Use menos Ajustar y más para acomodar árboles grandes.",
        ],
        accent=BURGUNDY,
    )

    page_break(doc)

    # Copiar ramas
    add_section_heading(
        doc,
        11,
        "Copiar y pegar ramas del árbol",
        "Esta función evita volver a ingresar estructuras repetidas. La copia crea datos independientes en el nodo de destino.",
        accent=BURGUNDY,
    )
    add_picture(
        doc,
        "11-arbol-copiar-ramas.png",
        "Figura 10  Controles para copiar ramas y estructura",
        "Editor de árbol con botones Copiar todas Copiar selección y Pegar estructura junto al diagrama",
        width=6.45,
    )
    add_steps(
        doc,
        [
            "Seleccione el nodo que contiene las ramas que desea reutilizar.",
            "Pulse Copiar todas para tomar la estructura completa marque varias ramas y use Copiar selección o use Copiar en una rama individual.",
            "Seleccione el nodo de destino en el diagrama.",
            "Pulse Pegar estructura. Las ramas actuales del destino serán reemplazadas por la copia.",
            "Edite nombres probabilidades y valores del destino. Los cambios no alteran el nodo original. Resuelva nuevamente el árbol.",
        ],
        accent=BURGUNDY,
    )

    page_break(doc)

    # Cierre y solución de problemas
    add_section_heading(
        doc,
        12,
        "Comprobaciones y solución de problemas",
        "Antes de comparar el resultado con un desarrollo manual o con POM QM revise las unidades los redondeos y el tipo de modelo.",
    )
    doc.add_heading("Comprobaciones rápidas", level=2)
    add_steps(
        doc,
        [
            "PERT CPM  Cada predecesora debe existir y aparecer antes de la actividad que depende de ella. Evite ciclos.",
            "Inventarios  Use datos del mismo periodo y confirme si el mantenimiento se ingresó por unidad o como porcentaje.",
            "Tablas de decisión  Verifique si la matriz contiene beneficios o costos y confirme que las probabilidades sumen 1 o 100.",
            "Árboles de decisión  Las probabilidades de las ramas de cada nodo de azar deben sumar 1 o 100.",
            "Después de editar datos pulse nuevamente el botón principal de cálculo para actualizar el resultado.",
        ],
    )
    doc.add_heading("Si no aparece el resultado", level=2)
    add_lead(doc, "Datos incompletos", "Revise los campos señalados por el sistema y complete valores nombres o probabilidades faltantes.")
    add_lead(doc, "Vista desordenada", "Pulse Ajustar en el diagrama. Después use los controles de zoom si necesita leer un nodo concreto.")
    add_lead(doc, "Datos anteriores", "Use Limpiar o Nueva matriz Nuevo árbol según el módulo. El guardado local puede restaurar el último ejercicio del navegador.")
    add_lead(doc, "Comparación de resultados", "El sistema conserva precisión interna y muestra resultados redondeados. Compare primero el resultado sin redondear y luego el valor mostrado.")
    doc.add_heading("Acceso directo", level=2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    add_hyperlink(p, "Abrir la plataforma de Investigación Operativa II", "https://io2modelos.vercel.app/")
    add_lead(doc, "Versión del manual", "14 de octubre de 2026.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
