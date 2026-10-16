from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
MANUAL_DIR = ROOT / "docs" / "manual"
SHOTS = MANUAL_DIR / "recursos"
OUT = MANUAL_DIR / "Manual_de_usuario_Investigacion_Operativa_II.docx"

NAVY = "0C2340"
GOLD = "C99A2E"
GREEN = "1B4D43"
BURGUNDY = "4A2238"
CORAL = "B9573F"
INK = "1B2430"
BODY = "3F4854"
MUTED = "6D7480"
PALE_BLUE = "F2F5F8"
PALE_GOLD = "FBF7EA"
BORDER = "D9D9D9"
WHITE = "FFFFFF"


def set_run_font(run, name="Aptos", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_keep(paragraph, keep_next=False, keep_together=True):
    paragraph.paragraph_format.keep_with_next = keep_next
    paragraph.paragraph_format.keep_together = keep_together


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def add_hyperlink(paragraph, text, url, color=NAVY):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Aptos")
    fonts.set(qn("w:hAnsi"), "Aptos")
    r_pr.extend([fonts, color_el, underline])
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color=BORDER, size=5):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:color"), color)
        borders.append(element)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def add_image_border(inline_shape, color="C7CDD4", width="9000"):
    sp_pr = inline_shape._inline.graphic.graphicData.pic.spPr
    line = OxmlElement("a:ln")
    line.set("w", width)
    solid = OxmlElement("a:solidFill")
    rgb = OxmlElement("a:srgbClr")
    rgb.set("val", color)
    solid.append(rgb)
    line.append(solid)
    sp_pr.append(line)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.14

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")
    title.paragraph_format.space_after = Pt(8)
    p_pr = title._element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)

    settings = {
        "Heading 1": (20, 13, 6),
        "Heading 2": (13.5, 10, 4),
        "Heading 3": (11.2, 7, 3),
    }
    for style_name, (size, before, after) in settings.items():
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = doc.styles["Caption"]
    caption.font.name = "Aptos"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(INK)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    add_page_number(p)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    r = p.add_run("UNJBG  |  Ingeniería en Informática y Sistemas  |  Octubre de 2026")
    set_run_font(r, size=8, color=MUTED)


def page_break(doc):
    doc.add_page_break()


def add_kicker(doc, text, color=NAVY):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    set_keep(p, keep_next=True)
    r = p.add_run(text.upper())
    set_run_font(r, size=8.5, color=color, bold=True)


def add_section(doc, number, title, intro, color=NAVY):
    add_kicker(doc, f"{number:02d}  Manual de usuario", color)
    h = doc.add_heading(title, level=1)
    h.paragraph_format.space_before = Pt(0)
    if intro:
        p = doc.add_paragraph(intro)
        p.paragraph_format.space_after = Pt(7)
        set_keep(p)


def add_figure(doc, number, title, filename, alt, width=6.25):
    label = doc.add_paragraph(style="Caption")
    label.paragraph_format.space_before = Pt(2)
    label.paragraph_format.space_after = Pt(0)
    set_keep(label, keep_next=True)
    r = label.add_run(f"Figura {number}")
    set_run_font(r, size=9, color=INK, bold=True)

    fig_title = doc.add_paragraph(style="Caption")
    fig_title.paragraph_format.space_before = Pt(0)
    fig_title.paragraph_format.space_after = Pt(4)
    set_keep(fig_title, keep_next=True)
    r = fig_title.add_run(title)
    set_run_font(r, size=9, color=INK, italic=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    set_keep(p)
    shape = p.add_run().add_picture(str(SHOTS / filename), width=Inches(width))
    add_image_border(shape)
    shape._inline.docPr.set("descr", alt)
    shape._inline.docPr.set("title", title)

    note = doc.add_paragraph(style="Caption")
    note.paragraph_format.space_before = Pt(0)
    note.paragraph_format.space_after = Pt(7)
    r = note.add_run("Nota. Captura de la plataforma elaborada por el equipo del proyecto.")
    set_run_font(r, size=8.3, color=MUTED, italic=True)


def add_steps(doc, steps, color=NAVY, heading="Procedimiento"):
    doc.add_heading(heading, level=2)
    for idx, item in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.31)
        p.paragraph_format.first_line_indent = Inches(-0.31)
        p.paragraph_format.space_after = Pt(4)
        set_keep(p)
        r = p.add_run(f"{idx:02d}  ")
        set_run_font(r, size=9.8, color=color, bold=True)
        if isinstance(item, tuple):
            label, body = item
            r = p.add_run(label + ". ")
            set_run_font(r, size=10.2, color=INK, bold=True)
            r = p.add_run(body)
            set_run_font(r, size=10.2, color=BODY)
        else:
            r = p.add_run(item)
            set_run_font(r, size=10.2, color=BODY)


def add_result(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Resultado esperado. ")
    set_run_font(r, size=10.2, color=INK, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.2, color=BODY)


def add_lead(doc, label, text, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(label + ". ")
    set_run_font(r, size=10.2, color=INK, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.2, color=BODY)


def add_toc_entry(doc, title, page, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25 if level == 2 else 0)
    p.paragraph_format.space_after = Pt(4 if level == 1 else 2)
    tab = p.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.45), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    r = p.add_run(title)
    set_run_font(r, size=10.5 if level == 1 else 9.8, color=INK, bold=(level == 1))
    r = p.add_run(f"\t{page}")
    set_run_font(r, size=10.5 if level == 1 else 9.8, color=MUTED)


def add_module_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.25, 2.35, 2.85]
    headers = ["Módulo", "Cuándo utilizarlo", "Resultado principal"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_fill(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=9, color=WHITE, bold=True)
    repeat_header(table.rows[0])
    rows = [
        ("PERT CPM", "Actividades con duración y precedencias", "Ruta crítica, red AON, Gantt y holguras"),
        ("Inventarios", "Reposición, producción o faltantes", "Lote óptimo, punto de reorden y costos"),
        ("Decisiones", "Alternativas bajo incertidumbre o riesgo", "Criterios, VME, EVPI y estrategia óptima"),
    ]
    for ridx, values in enumerate(rows, 1):
        row = table.add_row()
        for i, text in enumerate(values):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_fill(cell, WHITE if ridx % 2 else PALE_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell, 100, 105, 100, 105)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            set_run_font(r, size=9.2, color=BODY, bold=(i == 0))


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, size=10.2, color=BODY)


def build():
    doc = Document()
    configure_document(doc)

    # 1. Portada
    logos = doc.add_paragraph()
    logos.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logos.paragraph_format.space_after = Pt(12)
    r = logos.add_run()
    logo = r.add_picture(str(ROOT / "assets" / "institucional" / "logo-unjbg.png"), width=Inches(0.78))
    logo._inline.docPr.set("descr", "Escudo de la Universidad Nacional Jorge Basadre Grohmann")
    logo._inline.docPr.set("title", "Universidad Nacional Jorge Basadre Grohmann")
    logos.add_run("     ")
    r = logos.add_run()
    logo = r.add_picture(str(ROOT / "assets" / "institucional" / "logo-esis.png"), width=Inches(0.76))
    logo._inline.docPr.set("descr", "Logotipo de Ingeniería en Informática y Sistemas de la UNJBG")
    logo._inline.docPr.set("title", "Ingeniería en Informática y Sistemas")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Manual de usuario de la plataforma de Investigación Operativa II")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(11)
    r = subtitle.add_run("PERT CPM  |  Inventarios  |  Teoría de Decisiones")
    set_run_font(r, name="Aptos Display", size=14, color=INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(SHOTS / "01-portada.png"), width=Inches(5.8))
    add_image_border(shape, color=NAVY, width="11000")
    shape._inline.docPr.set("descr", "Página principal de la plataforma de Investigación Operativa II")
    shape._inline.docPr.set("title", "Vista principal de la plataforma")

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(8)
    meta.paragraph_format.space_after = Pt(4)
    r = meta.add_run("Universidad Nacional Jorge Basadre Grohmann\nEscuela Profesional de Ingeniería en Informática y Sistemas")
    set_run_font(r, size=10, color=INK, bold=True)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(4)
    r = meta.add_run("Investigación Operativa II  |  Docente MSc Manuel Yuri Apaza Valencia")
    set_run_font(r, size=9.2, color=MUTED)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("Tacna  2026")
    set_run_font(r, size=9.2, color=MUTED)

    page_break(doc)

    # 2. Presentación
    add_kicker(doc, "Información del manual")
    doc.add_heading("Presentación", level=1)
    add_lead(
        doc,
        "Propósito",
        "Orientar al estudiante en el ingreso de datos, la ejecución de los cálculos y la interpretación de los resultados de la plataforma.",
    )
    add_lead(
        doc,
        "Alcance",
        "El manual cubre PERT CPM, los modelos principales de inventarios, las tablas de decisión y los árboles de decisión. Incluye el flujo de trabajo, la lectura de resultados y las verificaciones previas a una comparación con POM QM.",
    )
    doc.add_heading("Requisitos de uso", level=2)
    add_bullets(
        doc,
        [
            "Navegador web actualizado con JavaScript habilitado.",
            "Conexión a Internet para abrir la versión publicada de la plataforma.",
            "Pantalla de computadora o tableta en orientación horizontal para visualizar mejor diagramas y tablas.",
            "Datos del ejercicio expresados en unidades y periodos compatibles.",
        ],
    )
    doc.add_heading("Cómo utilizar este manual", level=2)
    add_steps(
        doc,
        [
            ("Localice el módulo", "Use el contenido de la página siguiente para ir al procedimiento requerido."),
            ("Observe la figura", "Cada captura identifica la vista que debe aparecer en la plataforma."),
            ("Siga la secuencia", "Ejecute los pasos en orden y confirme el resultado esperado antes de continuar."),
        ],
        heading="",
    )
    doc.add_heading("Equipo responsable", level=2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("Elvis Ronald Arocutipa Quispe  ·  Helber Junior Lope Apomayta  ·  Antony Fernando Yucra Choquecota")
    set_run_font(r, size=9.6, color=BODY)
    p = doc.add_paragraph()
    r = p.add_run("Oscar Fabián Choque Palza  ·  William Andre Nina Estuco")
    set_run_font(r, size=9.6, color=BODY)

    page_break(doc)

    # 3. Índice estático
    add_kicker(doc, "Navegación")
    doc.add_heading("Contenido", level=1)
    p = doc.add_paragraph("Los procedimientos están ordenados según el recorrido normal de uso de la plataforma.")
    p.paragraph_format.space_after = Pt(12)
    entries = [
        ("Presentación", 2),
        ("Contenido", 3),
        ("1  Acceso y estructura general", 4),
        ("2  Ingreso de datos en PERT CPM", 5),
        ("3  Lectura de resultados en PERT CPM", 6),
        ("4  Desarrollo paso a paso en PERT CPM", 7),
        ("5  Guía teórica y modelamiento matemático", 8),
        ("6  Ingreso de datos en inventarios", 9),
        ("7  Lectura de resultados en inventarios", 10),
        ("8  Tablas de decisión", 11),
        ("9  Árboles de decisión", 12),
        ("10  Copiar y pegar ramas", 13),
        ("11  Validación y comparación de resultados", 14),
        ("12  Solución de problemas", 15),
    ]
    for text, number in entries:
        add_toc_entry(doc, text, number)
    doc.add_heading("Convenciones", level=2)
    add_lead(doc, "Botón principal", "Ejecuta o actualiza el cálculo del módulo activo.")
    add_lead(doc, "Caso de práctica", "Carga un conjunto completo de datos y reemplaza el formulario actual.")
    add_lead(doc, "Guía del método", "Abre los fundamentos, símbolos y fórmulas sin cerrar el ejercicio.")

    page_break(doc)

    # 4. Acceso
    add_section(
        doc,
        1,
        "Acceso y estructura general",
        "La plataforma organiza el trabajo en tres módulos. Todos conservan el mismo recorrido general: datos, cálculo y análisis.",
    )
    p = doc.add_paragraph()
    r = p.add_run("Acceso web. ")
    set_run_font(r, size=10.2, color=INK, bold=True)
    add_hyperlink(p, "Abrir la plataforma", "https://io2modelos.vercel.app/")
    add_figure(
        doc,
        1,
        "Selector de módulos de la plataforma",
        "02-modulos.png",
        "Tarjetas de acceso a PERT CPM Inventarios y Decisiones",
        width=6.1,
    )
    add_module_table(doc)
    add_steps(
        doc,
        [
            ("Abra Calculadoras", "También puede desplazarse hasta la sección de módulos desde la portada."),
            ("Elija el método", "Seleccione el módulo por el tipo de problema que desea resolver."),
            ("Cargue o escriba los datos", "Los casos de práctica permiten conocer el flujo antes de ingresar un ejercicio propio."),
        ],
    )
    add_result(doc, "El formulario del módulo elegido queda visible y listo para trabajar.")

    page_break(doc)

    # 5. PERT ingreso
    add_section(
        doc,
        2,
        "Ingreso de datos en PERT CPM",
        "Utilice CPM cuando cada actividad tenga una duración conocida. Utilice PERT cuando disponga de tiempos optimista, más probable y pesimista.",
        CORAL,
    )
    add_figure(
        doc,
        2,
        "Formulario de actividades PERT CPM",
        "03-pert-ingreso.png",
        "Formulario PERT CPM con selector de método tabla de actividades y botón Calcular proyecto",
        width=6.1,
    )
    add_steps(
        doc,
        [
            ("Seleccione el método", "Elija CPM determinístico o PERT de tres estimaciones."),
            ("Prepare la tabla", "Agregue filas, pegue datos desde Excel o cargue un archivo CSV."),
            ("Complete las actividades", "Ingrese identificador, descripción, predecesoras y duración. Separe varias predecesoras con comas, por ejemplo B, C."),
            ("Ejecute el cálculo", "Pulse Calcular proyecto cuando todas las filas estén completas."),
        ],
        CORAL,
    )
    add_result(doc, "La plataforma genera la red, la ruta crítica, el Gantt y el desarrollo matemático.")

    page_break(doc)

    # 6. PERT resultados
    add_section(
        doc,
        3,
        "Lectura de resultados en PERT CPM",
        "La red AON representa una actividad por nodo. El color crítico identifica las actividades que determinan la duración del proyecto.",
        CORAL,
    )
    add_figure(
        doc,
        3,
        "Red del proyecto y resumen de la ruta crítica",
        "04-pert-resultado.png",
        "Red AON con actividades ruta crítica controles de zoom y tarjetas de resultados",
        width=6.1,
    )
    add_steps(
        doc,
        [
            ("Ajuste el diagrama", "Use menos, Ajustar y más para cambiar el nivel de zoom."),
            ("Revise la ruta crítica", "Compruebe la secuencia, la duración total y la cantidad de actividades críticas."),
            ("Consulte un nodo", "Haga clic en una actividad para abrir su cálculo particular."),
            ("Cambie de vista", "Abra Gantt o Solución para revisar el cronograma y las operaciones."),
        ],
        CORAL,
    )
    add_result(doc, "La ruta crítica y la duración total coinciden con el recorrido sin holgura de la red.")

    page_break(doc)

    # 7. PERT desarrollo
    add_section(
        doc,
        4,
        "Desarrollo paso a paso en PERT CPM",
        "La vista Solución muestra las operaciones en el mismo orden de la red para facilitar la comparación con un procedimiento manual.",
        CORAL,
    )
    add_figure(
        doc,
        4,
        "Operaciones por actividad del proyecto",
        "05-pert-solucion.png",
        "Vista Solución con recorrido hacia adelante y operaciones por actividad",
        width=6.0,
    )
    add_steps(
        doc,
        [
            ("Avance por la red", "Tome el final más tardío de las predecesoras y sume la duración."),
            ("Retroceda desde el final", "Tome el inicio más próximo de las sucesoras y reste la duración."),
            ("Calcule la holgura", "Compare el inicio tardío con el inicio temprano. Un valor cero identifica una actividad crítica."),
            ("Relacione el cálculo", "Use Nodo para abrir la actividad equivalente en el diagrama."),
        ],
        CORAL,
    )
    add_result(doc, "Cada actividad muestra valores tempranos, tardíos y holgura verificables.")

    page_break(doc)

    # 8. Guía
    add_section(
        doc,
        5,
        "Guía teórica y modelamiento matemático",
        "La guía flotante permite consultar fundamentos, símbolos y fórmulas sin perder los datos del ejercicio actual.",
        CORAL,
    )
    add_figure(
        doc,
        5,
        "Fundamentos y formulación matemática de PERT CPM",
        "06-guia-metodo.png",
        "Ventana flotante con historia supuestos símbolos nodos y fórmulas PERT CPM",
        width=5.8,
    )
    add_steps(
        doc,
        [
            ("Abra la guía", "Pulse Guía del método desde cualquier vista del módulo."),
            ("Revise el fundamento", "Consulte el propósito, los supuestos y la diferencia entre CPM y PERT."),
            ("Interprete el nodo", "Relacione cada casilla con el inicio, la duración, el fin y la holgura."),
            ("Cierre la ventana", "Use la X para volver al análisis; el ejercicio permanece sin cambios."),
        ],
        CORAL,
    )
    add_result(doc, "El estudiante puede justificar el modelo y explicar las operaciones utilizadas.")

    page_break(doc)

    # 9. Inventarios ingreso
    add_section(
        doc,
        6,
        "Ingreso de datos en inventarios",
        "El módulo reúne EOQ con demanda constante, EPQ con producción gradual y EOQ con faltantes planificados.",
        GREEN,
    )
    add_figure(
        doc,
        6,
        "Formulario de parámetros de inventario",
        "07-inventarios-ingreso.png",
        "Formulario de inventarios con selección de modelo demanda costos días y moneda",
        width=6.1,
    )
    add_steps(
        doc,
        [
            ("Seleccione el modelo", "Elija EOQ, EPQ o faltantes según los datos del problema."),
            ("Use un mismo periodo", "Demanda, mantenimiento, producción y días deben corresponder al mismo horizonte."),
            ("Defina el mantenimiento", "Ingrese el costo por unidad al año o el porcentaje del precio, no ambos."),
            ("Cargue una práctica", "Seleccione un ejemplo y pulse Cargar para compararlo con el desarrollo del docente."),
            ("Calcule", "Pulse Calcular política óptima."),
        ],
        GREEN,
    )
    add_result(doc, "La política recomendada aparece junto con sus costos y su representación gráfica.")

    page_break(doc)

    # 10. Inventarios resultados
    add_section(
        doc,
        7,
        "Lectura de resultados en inventarios",
        "Los indicadores superiores resumen cuánto pedir o producir, cuándo reponer y con qué frecuencia hacerlo.",
        GREEN,
    )
    add_figure(
        doc,
        7,
        "Política óptima y evolución del inventario",
        "08-inventarios-resultado.png",
        "Resultados EOQ y gráfico de inventario con punto de reorden",
        width=6.1,
    )
    add_steps(
        doc,
        [
            ("Lea la política", "Identifique la cantidad óptima Q, el punto de reorden r, los pedidos por año N y el tiempo entre pedidos T."),
            ("Observe el gráfico", "La línea representa el nivel de existencias y la referencia horizontal marca el punto de pedido."),
            ("Revise los costos", "Abra Costos para separar ordenar, mantener, adquirir y otros componentes aplicables."),
            ("Compruebe el cálculo", "Abra Desarrollo y compare cada operación con el ejercicio manual."),
            ("Evalúe escenarios", "Use Comparar para observar el efecto de modificar un parámetro."),
        ],
        GREEN,
    )
    add_result(doc, "La recomendación escrita puede utilizarse directamente como política de abastecimiento.")

    page_break(doc)

    # 11. Tabla de decisión
    add_section(
        doc,
        8,
        "Tablas de decisión",
        "Esta herramienta compara alternativas frente a estados de la naturaleza mediante criterios bajo incertidumbre y riesgo.",
        BURGUNDY,
    )
    add_figure(
        doc,
        8,
        "Matriz de pagos y resultados por criterio",
        "09-decisiones-tabla.png",
        "Tabla de decisión con alternativas estados probabilidades y resumen de valor esperado",
        width=6.1,
    )
    add_steps(
        doc,
        [
            ("Defina el sentido", "Elija beneficios para seleccionar el mayor valor o costos para seleccionar el menor."),
            ("Complete la matriz", "Edite alternativas, estados y pagos directamente en la tabla."),
            ("Ingrese probabilidades", "La suma debe ser 1 o 100 por ciento para utilizar valor esperado y EVPI."),
            ("Ajuste Hurwicz", "Mueva el control de optimismo cuando corresponda."),
            ("Evalúe", "Pulse Evaluar decisiones y revise Criterios, Arrepentimiento y Desarrollo."),
        ],
        BURGUNDY,
    )
    add_result(doc, "La alternativa recomendada se presenta junto con el criterio que sustenta la elección.")

    page_break(doc)

    # 12. Árbol
    add_section(
        doc,
        9,
        "Árboles de decisión",
        "El árbol se utiliza cuando una decisión conduce a eventos posteriores. El sistema evalúa desde los resultados finales hasta el nodo inicial.",
        BURGUNDY,
    )
    add_figure(
        doc,
        9,
        "Editor y diagrama del árbol de decisión",
        "10-arbol-decision.png",
        "Árbol de decisión con editor de nodos controles de zoom y estrategia óptima",
        width=6.1,
    )
    add_steps(
        doc,
        [
            ("Inicie el caso", "Cargue una práctica o cree un árbol nuevo y defina si maximizará beneficios o minimizará costos."),
            ("Seleccione un nodo", "Edite su nombre, tipo y ramas desde el panel izquierdo."),
            ("Asigne el tipo", "Use Decisión para elegir, Azar para ponderar probabilidades y Resultado final para registrar un valor terminal."),
            ("Resuelva", "La estrategia óptima se resalta en dorado y el valor calculado aparece cerca de cada nodo."),
            ("Ajuste la vista", "Use menos, Ajustar y más para mantener legibles los árboles grandes."),
        ],
        BURGUNDY,
    )
    add_result(doc, "La primera decisión recomendada y el valor total del árbol quedan visibles en el resumen.")

    page_break(doc)

    # 13. Copiar ramas
    add_section(
        doc,
        10,
        "Copiar y pegar ramas",
        "La función permite reutilizar estructuras repetidas y después modificar la copia sin afectar el nodo original.",
        BURGUNDY,
    )
    add_figure(
        doc,
        10,
        "Controles para reutilizar una estructura de ramas",
        "11-arbol-copiar-ramas.png",
        "Editor del árbol con Copiar todas Copiar selección Copiar rama y Pegar estructura",
        width=6.1,
    )
    add_steps(
        doc,
        [
            ("Seleccione el origen", "Haga clic en el nodo que contiene las ramas que desea reutilizar."),
            ("Elija el alcance", "Use Copiar todas, marque varias ramas y pulse Copiar selección, o copie una sola rama."),
            ("Seleccione el destino", "El nodo debe admitir la estructura que se va a pegar."),
            ("Pegue la estructura", "Las ramas actuales del destino se reemplazan por una copia independiente."),
            ("Personalice y resuelva", "Cambie nombres, probabilidades o valores y pulse nuevamente Resolver árbol."),
        ],
        BURGUNDY,
    )
    add_result(doc, "La copia conserva nombres, probabilidades y nodos descendientes, pero puede editarse de forma independiente.")

    page_break(doc)

    # 14. Validación
    add_section(
        doc,
        11,
        "Validación y comparación de resultados",
        "Antes de comparar la plataforma con un desarrollo manual o con POM QM, confirme que ambos ejercicios utilizan los mismos datos y convenciones.",
    )
    doc.add_heading("Lista de verificación", level=2)
    add_steps(
        doc,
        [
            ("Modelo equivalente", "Seleccione en ambas herramientas el mismo método, por ejemplo CPM, EOQ o tabla de decisión."),
            ("Unidades compatibles", "No mezcle demanda mensual con costos anuales ni días calendario con días laborables."),
            ("Datos completos", "Compare actividad por actividad, parámetro por parámetro o celda por celda antes de resolver."),
            ("Criterio de redondeo", "La plataforma conserva precisión interna y redondea para mostrar; compare primero el valor sin redondear."),
            ("Interpretación", "Verifique no solo el número final, sino también la ruta, política o alternativa recomendada."),
        ],
    )
    doc.add_heading("Comprobaciones por módulo", level=2)
    add_lead(doc, "PERT CPM", "Las predecesoras deben existir, la red no debe contener ciclos y la ruta crítica debe tener holgura cero.")
    add_lead(doc, "Inventarios", "Confirme si el mantenimiento se expresa por unidad o como porcentaje del precio y revise el periodo de la demanda.")
    add_lead(doc, "Tablas de decisión", "Compruebe si la matriz contiene beneficios o costos y que las probabilidades sumen 1 o 100.")
    add_lead(doc, "Árboles de decisión", "Cada nodo de azar debe tener probabilidades completas y cada recorrido debe terminar en un resultado.")
    doc.add_heading("Registro recomendado", level=2)
    add_bullets(
        doc,
        [
            "Conserve una captura de los datos ingresados.",
            "Imprima o exporte el informe calculado.",
            "Anote cualquier diferencia causada por redondeos o convenciones del software comparado.",
        ],
    )

    page_break(doc)

    # 15. Problemas
    add_section(
        doc,
        12,
        "Solución de problemas",
        "La mayoría de incidencias se resuelve revisando campos incompletos, unidades, probabilidades o la escala del diagrama.",
    )
    doc.add_heading("El resultado no aparece", level=2)
    add_lead(doc, "Campos incompletos", "Revise los mensajes del formulario y complete nombres, valores, predecesoras o probabilidades faltantes.")
    add_lead(doc, "Probabilidades inválidas", "En decisiones y nodos de azar, la suma debe ser 1 o 100 por ciento.")
    add_lead(doc, "Cambios pendientes", "Después de modificar datos, pulse nuevamente el botón principal de cálculo.")
    doc.add_heading("La vista es difícil de leer", level=2)
    add_lead(doc, "Diagrama pequeño", "Pulse Ajustar y luego use los controles de zoom. En el navegador también puede reducir temporalmente la escala a 90 por ciento.")
    add_lead(doc, "Texto superpuesto", "Vuelva a pulsar Ajustar después de agregar o eliminar nodos y amplíe el ancho de la ventana.")
    add_lead(doc, "Datos anteriores", "Use Limpiar, Nueva matriz o Nuevo árbol. El guardado local puede restaurar el último ejercicio utilizado en ese navegador.")
    doc.add_heading("Preguntas frecuentes", level=2)
    add_lead(doc, "¿Puedo modificar un caso de práctica?", "Sí. Después de cargarlo, edite cualquier campo y ejecute nuevamente el cálculo.")
    add_lead(doc, "¿Dónde se guardan mis datos?", "En el almacenamiento local del navegador utilizado. No se sincronizan automáticamente con otros equipos.")
    add_lead(doc, "¿Por qué existe una pequeña diferencia con POM QM?", "Normalmente se debe al redondeo, al periodo seleccionado o a una convención distinta para los costos y probabilidades.")
    doc.add_heading("Acceso y versión", level=2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Plataforma. ")
    set_run_font(r, size=10.2, color=INK, bold=True)
    add_hyperlink(p, "Abrir Investigación Operativa II", "https://io2modelos.vercel.app/")
    add_lead(doc, "Versión del manual", "1.1, 14 de octubre de 2026.")
    add_lead(doc, "Uso académico", "Documento preparado para acompañar la demostración, el desarrollo manual y la comparación con POM QM.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
